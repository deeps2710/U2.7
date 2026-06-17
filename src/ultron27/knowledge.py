from __future__ import annotations

import hashlib
import json
import re
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Protocol


SUPPORTED_TEXT_SUFFIXES = {".md", ".markdown", ".txt"}
SUPPORTED_SUFFIXES = SUPPORTED_TEXT_SUFFIXES | {".pdf", ".docx"}
SECRET_PATTERN = re.compile(
    r"\b(password|passwd|secret|api[_ -]?key|token|credential|private key|bearer\s+[a-z0-9._-]+)\b",
    re.IGNORECASE,
)


class EmbeddingProvider(Protocol):
    name: str

    def embed(self, text: str) -> list[float]:
        ...


class KeywordOnlyEmbeddingProvider:
    name = "keyword_only"

    def embed(self, text: str) -> list[float]:
        return []


@dataclass(frozen=True)
class KnowledgeChunk:
    chunk_id: str
    text: str
    index: int


@dataclass(frozen=True)
class KnowledgeDocument:
    doc_id: str
    title: str
    path: str
    suffix: str
    ingested_at: float
    chunk_count: int
    metadata: dict[str, Any] = field(default_factory=dict)
    chunks: list[KnowledgeChunk] = field(default_factory=list)

    def public_metadata(self) -> dict[str, Any]:
        return {
            "doc_id": self.doc_id,
            "title": self.title,
            "path": self.path,
            "suffix": self.suffix,
            "ingested_at": self.ingested_at,
            "chunk_count": self.chunk_count,
            "metadata": self.metadata,
        }


class KnowledgeBase:
    def __init__(
        self,
        index_path: Path,
        *,
        workspace: Path,
        safe_roots: tuple[Path, ...],
        embedding_provider: EmbeddingProvider | None = None,
    ) -> None:
        self.index_path = _resolve_under_workspace(index_path, workspace)
        self.workspace = workspace.resolve()
        self.safe_roots = tuple(_resolve_under_workspace(root, self.workspace) for root in (safe_roots or (workspace,)))
        self.embedding_provider = embedding_provider or KeywordOnlyEmbeddingProvider()

    def ingest(self, source_path: str | Path) -> dict[str, Any]:
        path = _resolve_input_path(source_path, self.workspace)
        if not self._inside_safe_root(path):
            return {"status": "blocked", "message": f"Knowledge source is outside safe roots: {source_path}"}
        if not path.exists() or not path.is_file():
            return {"status": "not_found", "message": f"Knowledge source not found: {source_path}"}
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            return {"status": "unsupported", "message": f"Unsupported knowledge file type: {path.suffix}"}

        text_result = extract_text(path)
        if text_result["status"] != "ok":
            return text_result
        text = str(text_result["text"]).strip()
        if not text:
            return {"status": "empty", "message": f"No searchable text found in: {path.name}"}
        if looks_sensitive(text):
            return {"status": "rejected_sensitive", "message": "Knowledge source appears to contain secrets and was not stored."}

        chunks = chunk_text(text)
        doc_id = _document_id(path, text)
        document = KnowledgeDocument(
            doc_id=doc_id,
            title=path.stem,
            path=str(path),
            suffix=path.suffix.lower(),
            ingested_at=time.time(),
            chunk_count=len(chunks),
            metadata={"size_bytes": path.stat().st_size, "embedding_provider": self.embedding_provider.name},
            chunks=[KnowledgeChunk(chunk_id=f"{doc_id}:{index}", text=chunk, index=index) for index, chunk in enumerate(chunks)],
        )
        data = self._read()
        documents = [doc for doc in data["documents"] if doc.get("path") != str(path)]
        documents.append(_document_to_dict(document))
        self._write({"documents": documents})
        return {"status": "ok", "message": f"Ingested {path.name}.", "document": document.public_metadata()}

    def search(self, query: str, *, limit: int = 5) -> dict[str, Any]:
        query = query.strip()
        if not query:
            return {"status": "empty", "query": query, "matches": []}
        terms = tokenize(query)
        matches: list[dict[str, Any]] = []
        for document in self._documents():
            for chunk in document.chunks:
                score = score_chunk(terms, chunk.text)
                if score <= 0:
                    continue
                matches.append(
                    {
                        "score": score,
                        "doc_id": document.doc_id,
                        "title": document.title,
                        "path": document.path,
                        "chunk_id": chunk.chunk_id,
                        "snippet": snippet(chunk.text, terms),
                    }
                )
        matches.sort(key=lambda item: item["score"], reverse=True)
        return {"status": "ok", "query": query, "matches": matches[: max(1, min(limit, 25))], "sources": self.sources()}

    def sources(self) -> list[dict[str, Any]]:
        return [document.public_metadata() for document in self._documents()]

    def _documents(self) -> list[KnowledgeDocument]:
        documents: list[KnowledgeDocument] = []
        for raw in self._read()["documents"]:
            chunks = [
                KnowledgeChunk(chunk_id=str(chunk.get("chunk_id", "")), text=str(chunk.get("text", "")), index=int(chunk.get("index", 0)))
                for chunk in raw.get("chunks", [])
                if isinstance(chunk, dict)
            ]
            documents.append(
                KnowledgeDocument(
                    doc_id=str(raw.get("doc_id", "")),
                    title=str(raw.get("title", "")),
                    path=str(raw.get("path", "")),
                    suffix=str(raw.get("suffix", "")),
                    ingested_at=float(raw.get("ingested_at", 0.0)),
                    chunk_count=int(raw.get("chunk_count", len(chunks))),
                    metadata=raw.get("metadata") if isinstance(raw.get("metadata"), dict) else {},
                    chunks=chunks,
                )
            )
        return documents

    def _inside_safe_root(self, path: Path) -> bool:
        resolved = path.resolve()
        return any(_is_relative_to(resolved, root) for root in self.safe_roots)

    def _read(self) -> dict[str, list[dict[str, Any]]]:
        if not self.index_path.exists():
            return {"documents": []}
        try:
            payload = json.loads(self.index_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return {"documents": []}
        documents = payload.get("documents", []) if isinstance(payload, dict) else []
        return {"documents": [item for item in documents if isinstance(item, dict)]}

    def _write(self, payload: dict[str, list[dict[str, Any]]]) -> None:
        self.index_path.parent.mkdir(parents=True, exist_ok=True)
        self.index_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=True), encoding="utf-8")


def extract_text(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in SUPPORTED_TEXT_SUFFIXES:
        return {"status": "ok", "text": path.read_text(encoding="utf-8", errors="ignore")}
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return {"status": "not_implemented", "message": "PDF ingestion needs pypdf."}
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return {"status": "ok", "text": text}
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError:
            return {"status": "not_implemented", "message": "DOCX ingestion needs python-docx."}
        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return {"status": "ok", "text": text}
    return {"status": "unsupported", "message": f"Unsupported knowledge file type: {suffix}"}


def chunk_text(text: str, *, max_chars: int = 850) -> list[str]:
    paragraphs = [part.strip() for part in re.split(r"\n{2,}", text) if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(current) + len(paragraph) + 2 <= max_chars:
            current = f"{current}\n\n{paragraph}".strip()
            continue
        if current:
            chunks.append(current)
        if len(paragraph) <= max_chars:
            current = paragraph
        else:
            chunks.extend(paragraph[i : i + max_chars] for i in range(0, len(paragraph), max_chars))
            current = ""
    if current:
        chunks.append(current)
    return chunks or [text[:max_chars]]


def tokenize(text: str) -> list[str]:
    return [item for item in re.findall(r"[a-z0-9_]+", text.lower()) if len(item) > 2]


def score_chunk(terms: list[str], text: str) -> int:
    lowered = text.lower()
    return sum(lowered.count(term) for term in terms)


def snippet(text: str, terms: list[str], *, width: int = 240) -> str:
    lowered = text.lower()
    positions = [lowered.find(term) for term in terms if lowered.find(term) >= 0]
    start = max(0, min(positions) - 60) if positions else 0
    clean = " ".join(text[start : start + width].split())
    return clean


def looks_sensitive(value: str) -> bool:
    return bool(SECRET_PATTERN.search(value))


def _document_id(path: Path, text: str) -> str:
    digest = hashlib.sha256(f"{path.resolve()}:{path.stat().st_mtime_ns}:{len(text)}".encode("utf-8")).hexdigest()
    return digest[:16]


def _document_to_dict(document: KnowledgeDocument) -> dict[str, Any]:
    payload = asdict(document)
    payload["chunks"] = [asdict(chunk) for chunk in document.chunks]
    return payload


def _resolve_input_path(path: str | Path, workspace: Path) -> Path:
    candidate = Path(path).expanduser()
    if candidate.is_absolute():
        return candidate.resolve()
    return (workspace / candidate).resolve()


def _resolve_under_workspace(path: Path, workspace: Path) -> Path:
    candidate = path.expanduser()
    if not candidate.is_absolute():
        candidate = workspace / candidate
    return candidate.resolve()


def _is_relative_to(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root.resolve())
    except ValueError:
        return False
    return True
