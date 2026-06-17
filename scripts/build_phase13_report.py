from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_13_Report.docx"
BORDER = "B8C2CC"
HEADER_FILL = "F2F4F7"
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(
        doc,
        "Executive Summary",
        [
            "Phase 13 turns ULTRON 2.7 into a beta-ready local Windows prototype. The project now has setup scripts, dependency and provider checks, a config wizard, a one-command launcher, diagnostics, a no-microphone demo, and CI-style verification.",
            "The safety architecture is unchanged. Setup, diagnostics, demos, and verification are operator tooling around the existing safe runtime; they do not grant raw shell access or bypass schema validation, policy gates, executor restrictions, or audit logging.",
        ],
    )
    bullets(
        doc,
        "Phase 13 Scope",
        [
            "Add Windows setup and local dependency checks.",
            "Add STT/TTS/wake/VAD provider health checks.",
            "Add a config wizard for workspace, safe roots, app aliases, voice providers, and dry-run mode.",
            "Add a launcher that starts the backend and UI together and handles port conflicts.",
            "Add a diagnostics API and diagnostics page.",
            "Add a no-microphone demo flow using typed commands and mock voice.",
            "Add a CI-style verification script.",
            "Update README, architecture docs, install guide, run guide, troubleshooting guide, and safety guide.",
        ],
    )
    table(
        doc,
        "What Was Implemented",
        ("Area", "File(s)", "Result"),
        [
            ("Diagnostics", "src/ultron27/diagnostics.py, src/ultron27/web_server.py", "Adds /api/diagnostics with backend, brain, runtime path, voice provider, safe-root, dependency, and recent-error readiness data."),
            ("Launcher", "scripts/launch_ultron.py, scripts/run_phase13_beta.py", "Starts backend and UI together, chooses the next open port when needed, and prints the main and diagnostics URLs."),
            ("Setup", "scripts/setup_ultron_windows.ps1, scripts/check_dependencies.py", "Creates a Windows venv, installs the project, writes safe defaults, and reports missing files or packages clearly."),
            ("Config", "scripts/config_wizard.py", "Generates ultron.config.json with workspace, safe roots, app aliases, voice provider choices, and dry-run defaults."),
            ("Provider checks", "scripts/check_providers.py", "Reports active and configured STT/TTS/wake/VAD providers and fallback behavior."),
            ("Demo", "scripts/demo_phase13.py", "Runs typed command, note creation, confirmation pause, and mock voice demos without a microphone."),
            ("Verification", "scripts/verify_phase13.py", "Runs tests, dataset evaluation, safety regression, and API smoke checks."),
            ("UI", "web/diagnostics.html, web/diagnostics.js, web/styles.css", "Adds a readable diagnostics page linked from the main interface."),
            ("Docs", "docs/install_guide.md, docs/local_run_guide.md, docs/troubleshooting_guide.md, docs/safety_model_guide.md", "Adds beta-ready operator guides."),
        ],
    )
    table(
        doc,
        "Beta Operator Workflow",
        ("Need", "Command", "Expected Outcome"),
        [
            ("Install", ".\\scripts\\setup_ultron_windows.ps1", "Creates .venv, installs package, writes config, checks dependencies."),
            ("Launch", "python scripts\\launch_ultron.py", "Starts the local backend and UI, then prints URLs."),
            ("Diagnose", "Open /diagnostics", "Shows backend, brain, voice, memory, audit, safe roots, and recent errors."),
            ("Demo", "python scripts\\demo_phase13.py", "Runs typed and mock voice demos without a microphone."),
            ("Verify", "python scripts\\verify_phase13.py", "Runs tests, dataset evaluation, safety regression, and API smoke checks."),
        ],
    )
    table(
        doc,
        "Verification Results",
        ("Check", "Command", "Result"),
        [
            ("Unit tests", "python -m unittest discover -s tests", "71 tests passed."),
            ("Compile check", "python -m py_compile Phase 13 modules and scripts", "Python modules and scripts compiled successfully."),
            ("Frontend syntax", "node --check web/app.js and web/diagnostics.js", "Frontend JavaScript parsed successfully."),
            ("Dependency check", "python scripts/check_dependencies.py --json", "Required paths and packages were available."),
            ("Provider check", "python scripts/check_providers.py --json", "Text/browser providers active with no warnings."),
            ("Demo flow", "python scripts/demo_phase13.py", "Typed, note creation, confirmation, and mock voice demos completed."),
            ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Total 285; intent 0.9965; tool 0.9965; argument exact match 0.9579; confirmation 1.0."),
            ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases; tool accuracy 1.0; policy action accuracy 1.0."),
            ("CI-style verifier", "python scripts/verify_phase13.py", "Passed tests, dataset, safety, and API smoke checks."),
        ],
    )
    section(
        doc,
        "Safety Impact",
        [
            "Phase 13 does not add a new execution authority. Diagnostics are read-only, the launcher only starts the existing local web server, the demo uses typed and mock voice inputs, and the verifier exercises the same public API routes.",
            "All user goals continue through ULTRON's safe chain: brain/planner, typed tool call, schema validation, policy gate, executor, audit log, and response.",
        ],
    )
    bullets(
        doc,
        "Recommended Next Step",
        [
            "Package a signed Windows shortcut or small desktop launcher around scripts/launch_ultron.py.",
            "Add a first-run checklist in the diagnostics page for optional local voice model downloads.",
            "Add exportable demo logs for project presentations and beta testing feedback.",
        ],
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 13 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 13 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    p = doc.add_paragraph()
    r = p.add_run("Beta-ready local prototype")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)

    table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", "Phase 13: Beta-Ready Local Prototype"),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", "Implemented and verified"),
        ],
        header=False,
        widths=(Inches(1.55), Inches(4.95)),
    )


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(
    doc: Document,
    heading: str,
    columns: tuple[str, ...],
    rows: list[tuple[str, ...]],
    header: bool = True,
    widths: tuple[Inches, ...] | None = None,
) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        if len(all_rows[0]) == 3:
            widths = (Inches(1.35), Inches(2.35), Inches(2.8))
        else:
            widths = tuple(Inches(6.5 / len(all_rows[0])) for _ in all_rows[0])
    set_table_width(t, widths)
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0, fill=HEADER_FILL if header and r_idx == 0 else None)


def set_table_width(table_obj, widths) -> None:
    tbl_pr = table_obj._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width.twips for width in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table_obj._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width.twips))
        grid.append(col)


def set_cell(cell, bold: bool = False, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)

    if fill:
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill)

    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(8.9)
            run.bold = bold
            if bold:
                run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


if __name__ == "__main__":
    main()
