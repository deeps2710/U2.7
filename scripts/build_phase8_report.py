from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_8_Report.docx"
BORDER = "B8C2CC"
HEADER_FILL = "F2F4F7"


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(
        doc,
        "Executive Summary",
        [
            "Phase 8 adds voice interaction to ULTRON 2.7. The web interface can listen through browser speech recognition when available, send transcripts through the local API, process commands with the existing brain/runtime, and speak responses through browser speech synthesis.",
            "Voice mode remains an adapter around the safe runtime. It does not bypass typed tool calls, schema validation, policy gates, executor restrictions, audit logging, or confirmation requirements.",
        ],
    )
    bullets(
        doc,
        "Phase 8 Scope",
        [
            "Add a backend voice module with STT/TTS provider abstractions.",
            "Add voice API endpoints for start, stop, transcribe, speak, and status.",
            "Add microphone, push-to-talk, mute, stop-speaking, confirmation, and mock voice controls.",
            "Use browser speech-recognition events plus empty-transcript rejection as the lightweight VAD path.",
            "Add transcript history covering user speech, understood command, selected tool, result, and spoken response.",
            "Keep typed commands and subtitles available as fallbacks.",
            "Add tests with mocked speech providers and safe confirmation behavior.",
        ],
    )
    table(
        doc,
        "What Was Implemented",
        ("Area", "File(s)", "Result"),
        [
            ("Voice runtime", "src/ultron27/voice.py", "VoiceSession, transcript records, provider protocols, confirmation phrase handling, browser TTS delegation, and mock TTS."),
            ("Local API", "src/ultron27/web_server.py", "Added /api/voice/start, /api/voice/stop, /api/voice/transcribe, /api/speak, and /api/voice/status."),
            ("Web interface", "web/index.html, web/styles.css, web/app.js", "Added voice controls, transcript history, browser speech recognition, browser speech synthesis, mute, and stop voice."),
            ("Launcher", "scripts/run_phase8_voice_ui.py", "Adds a Phase 8-named launcher for the same local interface server."),
            ("Smoke check", "scripts/phase8_voice_smoke.mjs", "Adds browser smoke coverage when Playwright's browser binary is available."),
            ("Docs", "README.md, docs/architecture.md, docs/phase8_voice_mode.md", "Documents the voice flow, endpoints, safety boundary, and run instructions."),
            ("Tests", "tests/test_pipeline.py", "Expanded to 42 tests including mocked STT/TTS and high-risk voice confirmation."),
        ],
    )
    table(
        doc,
        "Voice Pipeline",
        ("Stage", "Purpose", "Safety Property"),
        [
            ("Capture", "Browser speech recognition or mock voice provides text.", "Capture does not execute actions; empty speech is ignored."),
            ("Transcription API", "POST /api/voice/transcribe accepts transcript payloads.", "Wake word is stripped and empty speech is rejected."),
            ("Brain/runtime", "Transcript is processed as a normal goal.", "Planner, validation, policy, executor, and audit remain authoritative."),
            ("Confirmation", "High-risk tasks pause for yes confirm or clicked confirmation.", "Destructive actions remain dry-run or not implemented."),
            ("Speech output", "Response text is sent to /api/speak and browser speech synthesis.", "Mute and stop-speaking controls remain available."),
        ],
    )
    table(
        doc,
        "User Controls",
        ("Control", "Behavior", "Fallback"),
        [
            ("Mic On/Off", "Starts or stops browser speech recognition.", "Typed command input remains available."),
            ("Push-to-talk", "Switches between one-shot and continuous recognition behavior.", "Mock Voice can exercise the pipeline without a microphone."),
            ("Mute", "Prevents audible ULTRON responses.", "Subtitles still show the response."),
            ("Stop Voice", "Cancels browser speech synthesis and marks speaking stopped.", "The UI can return to listening or idle."),
            ("Confirm", "Sends yes confirm for pending high-risk tasks.", "Policy still controls whether execution is allowed."),
        ],
    )
    table(
        doc,
        "Verification Results",
        ("Check", "Command", "Result"),
        [
            ("Unit tests", "$env:PYTHONPATH='src'; python -m unittest discover -s tests", "42 tests passed."),
            ("Compile check", "python -m py_compile src/ultron27/voice.py src/ultron27/web_server.py scripts/run_phase8_voice_ui.py", "Voice module, server, and launcher compiled successfully."),
            ("HTTP smoke", "local ThreadingHTTPServer endpoint check", "Voice controls served; start, transcribe, speak, and history endpoints worked."),
            ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 1.0, tool 1.0, argument exact match 0.9614, confirmation 1.0."),
            ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases, tool accuracy 1.0, policy action accuracy 1.0."),
            ("Browser smoke", "node scripts/phase8_voice_smoke.mjs", "Script degrades cleanly when the Playwright browser binary is unavailable in the sandbox."),
        ],
    )
    bullets(
        doc,
        "Recommended Next Step",
        [
            "Plug a local STT provider such as faster-whisper or whisper.cpp into the VoiceSession provider interface.",
            "Add a local TTS provider such as Piper or pyttsx3 for fully offline voice identity.",
            "Add wake-word and stronger VAD once the microphone path is stable across target machines.",
        ],
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 8 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 8 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    p = doc.add_paragraph()
    r = p.add_run("Voice interaction, transcript history, and speech response layer")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)

    table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", "Phase 8: Voice Mode"),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", "Implemented and verified"),
        ],
        header=False,
        widths=(Inches(1.55), Inches(4.95)),
    )


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(
    doc: Document,
    heading: str,
    columns: tuple[str, ...],
    rows: list[tuple[str, ...]],
    header: bool = True,
    widths: tuple[Inches, ...] | None = None,
) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        if len(all_rows[0]) == 3:
            widths = (Inches(1.45), Inches(2.35), Inches(2.7))
        else:
            widths = tuple(Inches(6.5 / len(all_rows[0])) for _ in all_rows[0])
    set_table_width(t, widths)
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0, fill=HEADER_FILL if header and r_idx == 0 else None)


def set_table_width(table_obj, widths) -> None:
    tbl_pr = table_obj._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width.twips for width in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table_obj._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width.twips))
        grid.append(col)


def set_cell(cell, bold: bool = False, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)

    if fill:
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill)

    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(9.1)
            run.bold = bold
            if bold:
                run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


if __name__ == "__main__":
    main()
