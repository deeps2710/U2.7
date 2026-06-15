from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_5_Report.docx"
BORDER = "B8C2CC"


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(
        doc,
        "Executive Summary",
        [
            "Phase 5 completes the basic working prototype by adding an interactive assistant console on top of the existing ULTRON safety pipeline.",
            "The new console keeps the assistant text-first while making repeated commands, confirmation flows, and JSON inspection practical for demos and iterative testing.",
        ],
    )
    bullets(
        doc,
        "Phase 5 Scope",
        [
            "Create a reusable runtime object for repeated assistant commands.",
            "Add an interactive console entry point with concise human-readable responses.",
            "Keep the same planner, validator, policy, executor, and audit path for one-shot and interactive usage.",
            "Add confirmation, JSON inspection, help, and exit commands inside the console.",
            "Document the final prototype workflow and provide a Windows convenience launcher.",
        ],
    )
    table(
        doc,
        "What Was Implemented",
        ("Area", "File(s)", "Result"),
        [
            ("Reusable runtime", "src/ultron27/runtime.py", "Added RuntimeSettings and UltronAssistant so commands can be handled repeatedly without reloading the dataset."),
            ("Interactive console", "src/ultron27/console.py", "Added /help, /json, /yes, and /exit commands plus concise text response formatting."),
            ("CLI integration", "src/ultron27/cli.py", "Added --interactive/-i and --text while preserving JSON output as the default."),
            ("Launcher", "scripts/run_ultron_console.ps1", "Added a Windows PowerShell helper to start the console from the repo root."),
            ("Docs", "README.md, docs/phase5_interactive_prototype.md", "Documented the interactive prototype and command examples."),
            ("Tests", "tests/test_pipeline.py", "Expanded to 24 regression tests covering runtime, text output, and console command parsing."),
        ],
    )
    table(
        doc,
        "Console Commands",
        ("Command", "Purpose"),
        [
            ("/help", "Show available console commands."),
            ("/json <command>", "Run a command and print the full JSON payload."),
            ("/yes <command>", "Confirm a command that requires confirmation."),
            ("/exit", "Leave the interactive console."),
        ],
    )
    table(
        doc,
        "Verification Results",
        ("Check", "Command", "Result"),
        [
            ("Unit tests", "python -m unittest discover -s tests", "24 tests passed."),
            ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 1.0, tool 1.0, argument exact match 0.9614, confirmation 1.0."),
            ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases, tool accuracy 1.0, policy action accuracy 1.0."),
            ("Text-mode smoke test", 'python -m ultron27 "delete project_report.txt" --text --no-audit', "High-risk command returned confirmation_required with /yes guidance."),
        ],
    )
    bullets(
        doc,
        "Prototype Status",
        [
            "The project now has a working text-first assistant loop.",
            "The default behavior remains safe dry-run execution.",
            "Low-risk actions can be executed through the allowlisted executor when --execute is used.",
            "Destructive actions remain intentionally unimplemented in the MVP.",
            "The next major layer should be voice input and voice output on top of this stable console runtime.",
        ],
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 12, "1F4D78")]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 5 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 5 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    p = doc.add_paragraph()
    r = p.add_run("Interactive prototype and reusable assistant runtime")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)
    table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", "Phase 5: Interactive Prototype"),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", "Implemented and verified"),
        ],
        header=False,
    )


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc: Document, heading: str, columns: tuple[str, ...], rows: list[tuple[str, ...]], header: bool = True) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0)


def set_cell(cell, bold: bool = False) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders") or OxmlElement("w:tcBorders")
    if borders.getparent() is None:
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}")) or OxmlElement(f"w:{edge}")
        if border.getparent() is None:
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(9.5)
            run.bold = bold


if __name__ == "__main__":
    main()
