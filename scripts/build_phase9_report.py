from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_9_Report.docx"
BORDER = "B8C2CC"
HEADER_FILL = "F2F4F7"


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(
        doc,
        "Executive Summary",
        [
            "Phase 9 upgrades ULTRON 2.7 from browser-only voice behavior to an offline-first provider stack. The system can now select local STT/TTS adapters, report provider health, fall back gracefully, and keep typed/browser operation available when local models are not installed.",
            "The safety architecture remains unchanged. Voice still flows through STT, the brain/runtime, typed tool calls, schema validation, policy gates, the executor, audit logging, response text, and TTS. No voice provider receives raw shell authority or executor authority.",
        ],
    )
    bullets(
        doc,
        "Phase 9 Scope",
        [
            "Add local STT provider adapters for faster-whisper and whisper.cpp.",
            "Add local TTS provider adapters for Piper and pyttsx3.",
            "Keep browser transcript and browser speech synthesis as fallbacks.",
            "Add voice provider config fields, environment overrides, and provider health checks.",
            "Expose /api/voice/providers, /api/voice/test-stt, and /api/voice/test-tts.",
            "Update the Phase 7/8 UI to display active STT/TTS providers and run provider tests.",
            "Add mocked provider-selection and health endpoint tests.",
        ],
    )
    table(
        doc,
        "What Was Implemented",
        ("Area", "File(s)", "Result"),
        [
            ("Configuration", "src/ultron27/config.py, ultron.config.example.json", "Added STT/TTS provider, model path, device, identity, rate, pitch, and volume settings."),
            ("Voice runtime", "src/ultron27/voice.py", "Added provider config, health model, faster-whisper, whisper.cpp, Piper, pyttsx3, browser, and mock adapters."),
            ("Local API", "src/ultron27/web_server.py", "Added provider health and test routes while preserving command and voice routes."),
            ("Web interface", "web/index.html, web/styles.css, web/app.js", "Added active provider display, refresh/test controls, and backend-aware browser speech fallback."),
            ("Launcher", "scripts/run_phase9_voice_ui.py", "Adds a Phase 9-named launcher for the local visual and voice interface."),
            ("Docs", "README.md, docs/architecture.md, docs/phase9_offline_voice_stack.md", "Documents configuration, provider matrix, health checks, and fallback behavior."),
            ("Tests", "tests/test_pipeline.py", "Expanded to 47 tests covering config parsing, provider fallback, mocked STT/TTS, and provider API endpoints."),
        ],
    )
    table(
        doc,
        "Provider Matrix",
        ("Provider", "Type", "Behavior"),
        [
            ("browser", "STT", "Uses browser speech recognition and sends transcript text to the backend."),
            ("text_payload", "STT", "Accepts transcript/text JSON payloads for deterministic tests and fallback runs."),
            ("faster_whisper", "STT", "Checks for the faster-whisper package and configured model path before activation."),
            ("whisper_cpp", "STT", "Checks for a whisper.cpp executable and configured model path before activation."),
            ("browser_speech_synthesis", "TTS", "Delegates audio playback to the browser client."),
            ("piper", "TTS", "Checks for Piper and a local voice model before activation."),
            ("pyttsx3", "TTS", "Uses the local pyttsx3 package when available."),
            ("mock", "STT/TTS", "Deterministic provider for tests and demos."),
        ],
    )
    table(
        doc,
        "Safety Boundary",
        ("Boundary", "Phase 9 Behavior", "Safety Property"),
        [
            ("STT", "Transcribes or accepts text payloads.", "Output is treated as a user goal, not as executable authority."),
            ("Brain/runtime", "Routes every goal through the existing task/runtime system.", "Typed tool calls, schema validation, policy, executor, and audit remain authoritative."),
            ("High-risk commands", "Still pause for explicit confirmation.", "Destructive actions remain dry-run or not implemented."),
            ("TTS", "Speaks or delegates the response text.", "Speech output cannot trigger tools."),
            ("Fallbacks", "Missing providers downgrade to browser/typed behavior.", "Unavailable local models do not crash the app."),
        ],
    )
    table(
        doc,
        "Verification Results",
        ("Check", "Command", "Result"),
        [
            ("Unit tests", "$env:PYTHONPATH='src'; python -m unittest discover -s tests", "47 tests passed."),
            ("Compile check", "python -m py_compile src/ultron27/config.py src/ultron27/voice.py src/ultron27/web_server.py scripts/run_phase9_voice_ui.py", "Phase 9 Python modules compiled successfully."),
            ("Frontend syntax", "bundled node --check web/app.js", "Frontend JavaScript parsed successfully."),
            ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 1.0, tool 1.0, argument exact match 0.9614, confirmation 1.0."),
            ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases, tool accuracy 1.0, policy action accuracy 1.0."),
            ("Provider API", "ThreadingHTTPServer endpoint tests", "Provider status, test-STT, and test-TTS routes returned expected payloads."),
        ],
    )
    bullets(
        doc,
        "Recommended Next Step",
        [
            "Install a real faster-whisper or whisper.cpp model and verify microphone audio capture end to end.",
            "Install a Piper voice model or pyttsx3 on the target Windows machine and tune ULTRON's voice identity.",
            "Add stronger wake-word and VAD layers after the offline providers are stable.",
        ],
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 9 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 9 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    p = doc.add_paragraph()
    r = p.add_run("Offline-first voice provider stack")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)

    table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", "Phase 9: Offline Voice Stack"),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", "Implemented and verified"),
        ],
        header=False,
        widths=(Inches(1.55), Inches(4.95)),
    )


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(
    doc: Document,
    heading: str,
    columns: tuple[str, ...],
    rows: list[tuple[str, ...]],
    header: bool = True,
    widths: tuple[Inches, ...] | None = None,
) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        if len(all_rows[0]) == 3:
            widths = (Inches(1.35), Inches(2.35), Inches(2.8))
        else:
            widths = tuple(Inches(6.5 / len(all_rows[0])) for _ in all_rows[0])
    set_table_width(t, widths)
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0, fill=HEADER_FILL if header and r_idx == 0 else None)


def set_table_width(table_obj, widths) -> None:
    tbl_pr = table_obj._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width.twips for width in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table_obj._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width.twips))
        grid.append(col)


def set_cell(cell, bold: bool = False, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)

    if fill:
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill)

    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(9.1)
            run.bold = bold
            if bold:
                run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


if __name__ == "__main__":
    main()
