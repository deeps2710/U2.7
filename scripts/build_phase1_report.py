from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_1_Report.docx"
WORKSPACE_LABEL = "Local workspace checkout: U2.7"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
GRAY_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
BORDER = "B8C2CC"


def main() -> None:
    doc = Document()
    configure_document(doc)

    add_title(doc)
    add_executive_summary(doc)
    add_phase_scope(doc)
    add_architecture(doc)
    add_implementation_table(doc)
    add_safety_model(doc)
    add_verification(doc)
    add_repository_map(doc)
    add_handoff(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("ULTRON 2.7 Phase 1 Report")
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_title(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("ULTRON 2.7 Phase 1 Execution Report")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Basic text-first prototype for a safe Jarvis-style laptop assistant")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    meta = doc.add_table(rows=4, cols=2)
    style_table(meta, [2200, 7160], fill_header=False)
    rows = [
        ("Project", "ULTRON 2.7"),
        ("Phase", "Phase 1: Core Text Prototype"),
        ("Workspace", WORKSPACE_LABEL),
        ("Status", "Implemented and verified"),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True


def add_executive_summary(doc: Document) -> None:
    doc.add_heading("Executive Summary", level=1)
    add_para(
        doc,
        "Phase 1 has been executed as a safe text-first assistant prototype. "
        "The implementation converts a user command into a structured tool call, validates the call against an allowlisted schema, applies risk and confirmation policy, runs only through a controlled executor, and writes an audit record.",
    )
    add_para(
        doc,
        "The prototype intentionally defaults to dry-run execution. That means the assistant can be tested against real commands without making risky changes to the laptop. High-risk commands, such as deleting files, stop at a confirmation gate.",
    )


def add_phase_scope(doc: Document) -> None:
    doc.add_heading("Phase 1 Scope", level=1)
    add_bullets(
        doc,
        [
            "Build a command-line assistant entry point.",
            "Use the supplied dataset for exact and fuzzy command planning.",
            "Add deterministic regex fallbacks for common numeric commands.",
            "Define typed tool schemas and argument validation.",
            "Add policy gates for low, medium, high, and unsupported actions.",
            "Add a dry-run executor and audit log.",
            "Add tests and dataset evaluation for the basic prototype.",
        ],
    )


def add_architecture(doc: Document) -> None:
    doc.add_heading("Implemented Architecture", level=1)
    add_para(doc, "The current Phase 1 pipeline is:")
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.25)
    run = code.add_run("text command -> planner -> tool call -> validator -> policy -> executor -> audit log")
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    shade_paragraph(code, "F4F6F9")
    add_para(
        doc,
        "The design follows the research paper's recommendation that the LLM or planner should propose actions, while the application remains responsible for validation, permission checks, execution, and logging.",
    )


def add_implementation_table(doc: Document) -> None:
    doc.add_heading("What Was Implemented", level=1)
    table = doc.add_table(rows=1, cols=3)
    style_table(table, [2500, 4000, 2860])
    set_row(table.rows[0], ["Component", "File(s)", "Purpose"], header=True)
    rows = [
        ("CLI", "src/ultron27/cli.py", "Accepts a typed command, prints JSON output, and writes audit records."),
        ("Planner", "src/ultron27/planner.py", "Uses dataset matching plus deterministic fallbacks to produce a tool call."),
        ("Tool schemas", "src/ultron27/tools.py", "Defines allowed tools, required arguments, numeric ranges, and validation rules."),
        ("Policy", "src/ultron27/policy.py", "Allows safe commands, confirms risky commands, and blocks unsupported requests."),
        ("Executor", "src/ultron27/executor.py", "Runs dry-run actions by default and contains safe handlers for early tools."),
        ("Evaluation", "src/ultron27/evaluate.py, scripts/evaluate_dataset.py", "Measures planner quality against the supplied dataset splits."),
        ("Tests", "tests/test_pipeline.py", "Covers parsing, validation, policy, and dry-run execution behavior."),
        ("Docs", "README.md, docs/architecture.md, docs/evaluation.md", "Explains usage, architecture, evaluation, and next steps."),
    ]
    for row in rows:
        set_row(table.add_row(), row)


def add_safety_model(doc: Document) -> None:
    doc.add_heading("Safety Model", level=1)
    add_para(doc, "Phase 1 enforces a tool-based safety model instead of direct shell access.")
    table = doc.add_table(rows=1, cols=3)
    style_table(table, [1700, 3600, 4060])
    set_row(table.rows[0], ["Risk", "Examples", "Prototype behavior"], header=True)
    rows = [
        ("None", "clarification, unsupported request", "No OS action is executed."),
        ("Low", "open app, set volume, search files", "Allowed after schema validation."),
        ("Medium", "draft email, move file, calendar draft", "Confirmation when configured by the tool."),
        ("High", "delete file, send email, run script, shutdown", "Confirmation required; destructive execution is not implemented."),
    ]
    for row in rows:
        set_row(table.add_row(), row)


def add_verification(doc: Document) -> None:
    doc.add_heading("Verification Results", level=1)
    add_para(doc, "The Phase 1 prototype was verified with unit tests, dataset evaluation, and a high-risk CLI safety check.")
    table = doc.add_table(rows=1, cols=3)
    style_table(table, [2800, 3000, 3560])
    set_row(table.rows[0], ["Check", "Command", "Result"], header=True)
    rows = [
        ("Unit tests", "python -m unittest discover -s tests", "6 tests passed."),
        ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 1.0, tool 1.0, argument exact match 0.9614, confirmation 1.0."),
        ("High-risk command", 'python -m ultron27 "delete project_report.txt"', "Policy returned confirmation_required."),
    ]
    for row in rows:
        set_row(table.add_row(), row)
    add_para(
        doc,
        "The argument exact-match misses come from ambiguous synthetic dataset rows where similar or identical utterances have different generated file types, folders, or file names. That is documented in docs/evaluation.md and should be cleaned during dataset improvement.",
    )


def add_repository_map(doc: Document) -> None:
    doc.add_heading("Repository Map", level=1)
    add_bullets(
        doc,
        [
            "src/ultron27: assistant package and core Phase 1 logic.",
            "scripts/evaluate_dataset.py: dataset evaluation entry point.",
            "tests/test_pipeline.py: regression tests for command planning and safety gates.",
            "data/jarvis_dataset_v2: supplied synthetic command dataset.",
            "docs/jarvis_research_paper.md: supplied research paper.",
            "docs/architecture.md: local-first architecture summary.",
            "docs/evaluation.md: baseline metrics and dataset notes.",
        ],
    )


def add_handoff(doc: Document) -> None:
    doc.add_heading("How To Run Phase 1", level=1)
    add_numbered(
        doc,
        [
            "Set PYTHONPATH to src when running from the checkout, or install with pip install -e .",
            'Run a safe command: python -m ultron27 "set volume to 40 percent".',
            'Check a high-risk gate: python -m ultron27 "delete project_report.txt".',
            "Run tests: python -m unittest discover -s tests.",
            "Run evaluation: python scripts/evaluate_dataset.py --split test.",
        ],
    )
    doc.add_heading("Recommended Next Phase", level=1)
    add_para(
        doc,
        "Phase 2 should focus on real low-risk laptop execution: app opening, file search, folder opening, note creation, reminders, clipboard, screenshot, volume, and brightness adapters. High-risk actions should remain preview-only until the confirmation UX is stronger.",
    )


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def style_table(table, widths: list[int], fill_header: bool = True) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_tbl_width(table, sum(widths), 120)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[col_index])
            set_cell_margins(cell)
            set_cell_border(cell)
            if fill_header and row_index == 0:
                shade_cell(cell, GRAY_FILL)


def set_row(row, values, header: bool = False) -> None:
    for cell, value in zip(row.cells, values):
        cell.text = str(value)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if header:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


def set_tbl_width(table, width: int, indent: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), BORDER)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


if __name__ == "__main__":
    main()
