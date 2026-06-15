from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_6_Report.docx"
BORDER = "B8C2CC"


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(
        doc,
        "Executive Summary",
        [
            "Phase 6 adds ULTRON's brain layer: a task-level orchestration system that can plan multi-step goals, execute allowed steps, pause for confirmation, remember non-sensitive context, and summarize outcomes.",
            "The brain does not receive raw shell access. Every step still passes through the existing planner, schema validator, policy gate, executor, and audit log.",
        ],
    )
    bullets(
        doc,
        "Phase 6 Scope",
        [
            "Add task state models and a TaskPlan structure.",
            "Add UltronBrain for goal classification, decomposition, preview, execution, and summaries.",
            "Add short-term session memory and persistent non-sensitive memory.",
            "Add console and CLI commands for /plan, /do, /memory, and /forget.",
            "Keep LLM and agentic behavior inside the existing safe runtime boundary.",
        ],
    )
    table(
        doc,
        "What Was Implemented",
        ("Area", "File(s)", "Result"),
        [
            ("Brain layer", "src/ultron27/brain.py", "TaskState, TaskStep, TaskPlan, MemoryStore, and UltronBrain."),
            ("Planner support", "src/ultron27/planner.py", "Added deterministic append_to_note parsing for multi-step note workflows."),
            ("Console commands", "src/ultron27/console.py", "Added /plan, /do, /memory, and /forget."),
            ("CLI commands", "src/ultron27/cli.py", "Slash commands now work in one-shot mode and interactive mode."),
            ("Docs", "README.md, docs/architecture.md, docs/phase6_brain_layer.md", "Documented the brain boundary, memory, and command examples."),
            ("Tests", "tests/test_pipeline.py", "Expanded to 31 tests covering brain behavior and safety boundaries."),
        ],
    )
    table(
        doc,
        "Brain Workflow",
        ("Stage", "Purpose", "Safety Note"),
        [
            ("Goal intake", "Receive a user goal and classify single-step vs multi-step.", "No execution happens during classification."),
            ("TaskPlan", "Create ordered TaskStep entries with previewed tool calls.", "Preview uses planner, validator, and policy."),
            ("Execution", "Send each step to UltronAssistant.handle().", "Runtime still enforces validation, policy, executor, and audit."),
            ("Pause/stop", "Stop on confirmation, blocked, or failed steps.", "High-risk actions cannot silently continue."),
            ("Memory", "Record small non-sensitive facts and completed-task summaries.", "Sensitive strings such as passwords and API keys are rejected."),
        ],
    )
    table(
        doc,
        "Verification Results",
        ("Check", "Command", "Result"),
        [
            ("Unit tests", "python -m unittest discover -s tests", "31 tests passed."),
            ("Acceptance plan", "python -m ultron27 /plan ... --text --no-audit", "Goal decomposed into create_note and append_to_note."),
            ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 1.0, tool 1.0, argument exact match 0.9614, confirmation 1.0."),
            ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases, tool accuracy 1.0, policy action accuracy 1.0."),
        ],
    )
    bullets(
        doc,
        "Recommended Next Step",
        [
            "Move to Phase 7 by connecting this brain runtime to a dedicated visual interface.",
            "Keep the interface as a client of the brain/runtime API rather than duplicating planner logic.",
            "Preserve the same confirmation and memory boundaries before adding voice in Phase 8.",
        ],
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 12, "1F4D78")]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 6 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 6 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    p = doc.add_paragraph()
    r = p.add_run("Agentic brain layer and safe multi-step task runtime")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)
    table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", "Phase 6: Brain Layer"),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", "Implemented and verified"),
        ],
        header=False,
    )


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc: Document, heading: str, columns: tuple[str, ...], rows: list[tuple[str, ...]], header: bool = True) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0)


def set_cell(cell, bold: bool = False) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders") or OxmlElement("w:tcBorders")
    if borders.getparent() is None:
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}")) or OxmlElement(f"w:{edge}")
        if border.getparent() is None:
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(9.5)
            run.bold = bold


if __name__ == "__main__":
    main()
