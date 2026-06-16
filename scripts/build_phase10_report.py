from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_10_Report.docx"
BORDER = "B8C2CC"
HEADER_FILL = "F2F4F7"


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(
        doc,
        "Executive Summary",
        [
            "Phase 10 adds wake-word detection and stronger voice activity detection to ULTRON 2.7. ULTRON can now enter an always-listening mode, wait for ULTRON or Hey ULTRON, ignore empty/noisy input, and only send authorized speech segments into STT and the safe runtime.",
            "Wake word and VAD are input gates only. They cannot choose tools, execute commands, skip validation, bypass policy, or bypass audit logging.",
        ],
    )
    bullets(
        doc,
        "Phase 10 Scope",
        [
            "Add wake-word and VAD provider abstractions.",
            "Add text/mock wake-word providers and health-checked openWakeWord support.",
            "Add energy-threshold VAD plus health-checked Silero and WebRTC adapters.",
            "Add always-listening modes: inactive, waiting_for_wake_word, listening, transcribing, thinking, and speaking.",
            "Add wake APIs and UI privacy controls.",
            "Keep push-to-talk and typed input available.",
            "Add regression tests for wake gating, noisy input, VAD-gated STT, and unchanged high-risk confirmation behavior.",
        ],
    )
    table(
        doc,
        "What Was Implemented",
        ("Area", "File(s)", "Result"),
        [
            ("Wake/VAD runtime", "src/ultron27/wake.py", "Provider protocols, text/openWakeWord/mock wake providers, energy/Silero/WebRTC/mock VAD providers, health checks, and wake phrase stripping."),
            ("Voice session", "src/ultron27/voice.py", "Always-listening state, wake snapshots, wake processing, VAD gating, and safe STT handoff."),
            ("Local API", "src/ultron27/web_server.py", "Added /api/wake/start, /api/wake/stop, /api/wake/status, and /api/wake/process."),
            ("Configuration", "src/ultron27/config.py, ultron.config.example.json", "Added wake provider, wake phrases, wake model path, VAD provider, and VAD threshold settings."),
            ("Web interface", "web/index.html, web/styles.css, web/app.js", "Added always-listening toggle, microphone privacy status, wake gate status, VAD status, and ignored-input feedback."),
            ("Launcher", "scripts/run_phase10_wake_ui.py", "Adds a Phase 10 launcher for the visual voice interface."),
            ("Docs", "README.md, docs/architecture.md, docs/phase10_wake_word_vad.md", "Documents wake/VAD flow, providers, APIs, and safety boundaries."),
            ("Tests", "tests/test_pipeline.py", "Expanded to 54 tests covering wake word, no-wake ignore, noisy input, VAD-gated STT, wake APIs, config parsing, and safety regression."),
        ],
    )
    table(
        doc,
        "Wake/VAD Pipeline",
        ("Stage", "Purpose", "Safety Property"),
        [
            ("Candidate input", "Browser or mock voice submits a candidate speech segment.", "No command execution happens at capture time."),
            ("VAD", "Empty, noisy, or low-energy input is ignored.", "Noisy audio does not reach STT or the brain."),
            ("Wake word", "ULTRON or Hey ULTRON must be present while waiting.", "Background speech without wake phrase is ignored."),
            ("STT", "Authorized speech is transcribed through the existing provider layer.", "STT output is still untrusted user input."),
            ("Brain/runtime", "The command flows through the same typed tool pipeline.", "Schema validation, policy, executor, and audit remain authoritative."),
        ],
    )
    table(
        doc,
        "Verification Results",
        ("Check", "Command", "Result"),
        [
            ("Unit tests", "$env:PYTHONPATH='src'; python -m unittest discover -s tests", "54 tests passed."),
            ("Compile check", "python -m py_compile src/ultron27/config.py src/ultron27/wake.py src/ultron27/voice.py src/ultron27/web_server.py scripts/run_phase10_wake_ui.py", "Phase 10 Python modules compiled successfully."),
            ("Frontend syntax", "bundled node --check web/app.js", "Frontend JavaScript parsed successfully."),
            ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 1.0, tool 1.0, argument exact match 0.9614, confirmation 1.0."),
            ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases, tool accuracy 1.0, policy action accuracy 1.0."),
            ("Wake API smoke", "ThreadingHTTPServer endpoint checks", "Wake start, status, process, and stop routes returned expected payloads."),
        ],
    )
    bullets(
        doc,
        "Recommended Next Step",
        [
            "Install real openWakeWord and Silero/WebRTC dependencies for microphone audio.",
            "Connect browser microphone chunks or local audio capture to the wake/VAD providers.",
            "Add streaming partial transcripts after the wake path is stable.",
        ],
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 10 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 10 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    p = doc.add_paragraph()
    r = p.add_run("Wake-word detection and voice activity gating")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)

    table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", "Phase 10: Wake Word and VAD"),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", "Implemented and verified"),
        ],
        header=False,
        widths=(Inches(1.55), Inches(4.95)),
    )


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(
    doc: Document,
    heading: str,
    columns: tuple[str, ...],
    rows: list[tuple[str, ...]],
    header: bool = True,
    widths: tuple[Inches, ...] | None = None,
) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        if len(all_rows[0]) == 3:
            widths = (Inches(1.35), Inches(2.35), Inches(2.8))
        else:
            widths = tuple(Inches(6.5 / len(all_rows[0])) for _ in all_rows[0])
    set_table_width(t, widths)
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0, fill=HEADER_FILL if header and r_idx == 0 else None)


def set_table_width(table_obj, widths) -> None:
    tbl_pr = table_obj._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width.twips for width in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table_obj._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width.twips))
        grid.append(col)


def set_cell(cell, bold: bool = False, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)

    if fill:
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill)

    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(9.1)
            run.bold = bold
            if bold:
                run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


if __name__ == "__main__":
    main()
