from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
BORDER = "B8C2CC"
HEADER_FILL = "F2F4F7"
INK = "0B2545"
MUTED = "555555"
HEADING = "2E74B5"
HEADING_DARK = "1F4D78"

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def build_docx(report: dict[str, Any], out: Path) -> None:
    doc = Document()
    configure_docx(doc, report)
    docx_title(doc, report)
    add_docx_content(doc, report)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out.name)


def configure_docx(doc: Document, report: dict[str, Any]) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, HEADING, 16, 8),
        ("Heading 2", 13, HEADING, 12, 6),
        ("Heading 3", 12, HEADING_DARK, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(f"ULTRON 2.7 Phase {report['phase']} Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def docx_title(doc: Document, report: dict[str, Any]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(report["title"])
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    r = p.add_run(report["subtitle"])
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    p.paragraph_format.space_after = Pt(14)

    docx_table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", report["phase_label"]),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", report["status"]),
        ],
        header=False,
        widths=(Inches(1.55), Inches(4.95)),
    )


def add_docx_content(doc: Document, report: dict[str, Any]) -> None:
    for block in report["blocks"]:
        kind = block["kind"]
        if kind == "section":
            doc.add_heading(block["heading"], level=1)
            for paragraph in block["paragraphs"]:
                doc.add_paragraph(paragraph)
        elif kind == "bullets":
            doc.add_heading(block["heading"], level=1)
            for item in block["items"]:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            docx_table(
                doc,
                block["heading"],
                tuple(block["columns"]),
                [tuple(row) for row in block["rows"]],
                header=True,
                widths=tuple(Inches(width) for width in block.get("widths", [1.35, 2.35, 2.8])),
            )
        else:
            raise ValueError(f"Unknown report block kind: {kind}")


def docx_table(
    doc: Document,
    heading: str,
    columns: tuple[str, ...],
    rows: list[tuple[str, ...]],
    header: bool = True,
    widths: tuple[Inches, ...] | None = None,
) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        widths = tuple(Inches(6.5 / len(all_rows[0])) for _ in all_rows[0])
    set_docx_table_width(t, widths)
    for row_index, row_values in enumerate(all_rows):
        for col_index, value in enumerate(row_values):
            cell = t.rows[row_index].cells[col_index]
            cell.text = str(value)
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_docx_cell(
                cell,
                bold=header and row_index == 0,
                fill=HEADER_FILL if header and row_index == 0 else None,
            )


def set_docx_table_width(table_obj, widths) -> None:
    tbl_pr = table_obj._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width.twips for width in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table_obj._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width.twips))
        grid.append(col)


def set_docx_cell(cell, bold: bool = False, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)

    if fill:
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill)

    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(8.9)
            run.bold = bold
            if bold:
                run.font.color.rgb = RGBColor.from_string(INK)


def build_pdf(report: dict[str, Any], out: Path) -> None:
    styles = pdf_stylesheet()
    doc = SimpleDocTemplate(
        str(out),
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )
    story: list[Any] = [
        Paragraph(report["title"], styles["TitleCustom"]),
        Paragraph(report["subtitle"], styles["SubtitleCustom"]),
        pdf_table(
            [
                ["Project", "ULTRON 2.7"],
                ["Phase", report["phase_label"]],
                ["Workspace", "Local workspace checkout: U2.7"],
                ["Status", report["status"]],
            ],
            [1.55 * inch, 4.95 * inch],
            styles,
        ),
        Spacer(1, 8),
    ]
    for block in report["blocks"]:
        kind = block["kind"]
        if kind == "section":
            story.append(Paragraph(block["heading"], styles["H1"]))
            story.extend(Paragraph(paragraph, styles["Normal"]) for paragraph in block["paragraphs"])
        elif kind == "bullets":
            story.append(Paragraph(block["heading"], styles["H1"]))
            story.append(pdf_bullets(block["items"], styles))
        elif kind == "table":
            story.append(Paragraph(block["heading"], styles["H1"]))
            widths = [width * inch for width in block.get("widths", [1.35, 2.35, 2.8])]
            story.append(pdf_table([block["columns"], *block["rows"]], widths, styles, header=True))
        else:
            raise ValueError(f"Unknown report block kind: {kind}")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.build(
        story,
        onFirstPage=lambda canvas, doc_obj: pdf_footer(canvas, doc_obj, report),
        onLaterPages=lambda canvas, doc_obj: pdf_footer(canvas, doc_obj, report),
    )
    print(out.name)


def pdf_stylesheet():
    styles = getSampleStyleSheet()
    styles["Normal"].fontName = "Helvetica"
    styles["Normal"].fontSize = 10.2
    styles["Normal"].leading = 12.8
    styles["Normal"].spaceAfter = 6
    styles.add(
        ParagraphStyle(
            "TitleCustom",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=29,
            textColor=colors.HexColor(f"#{INK}"),
            alignment=TA_LEFT,
            spaceAfter=3,
        )
    )
    styles.add(
        ParagraphStyle(
            "SubtitleCustom",
            parent=styles["Normal"],
            fontSize=12,
            leading=15,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=14,
        )
    )
    styles.add(
        ParagraphStyle(
            "H1",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=colors.HexColor(f"#{HEADING}"),
            spaceBefore=16,
            spaceAfter=8,
        )
    )
    styles.add(ParagraphStyle("Cell", parent=styles["Normal"], fontSize=7.95, leading=10.1, spaceAfter=0))
    styles.add(ParagraphStyle("CellBold", parent=styles["Cell"], fontName="Helvetica-Bold", textColor=colors.HexColor(f"#{INK}")))
    return styles


def pdf_table(rows, widths, styles, header: bool = False):
    converted = [
        [Paragraph(str(cell), styles["CellBold"] if header and row_index == 0 else styles["Cell"]) for cell in row]
        for row_index, row in enumerate(rows)
    ]
    table = Table(converted, colWidths=widths, hAlign="LEFT", repeatRows=1 if header else 0)
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor(f"#{BORDER}")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if header:
        commands.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{HEADER_FILL}")))
    table.setStyle(TableStyle(commands))
    return table


def pdf_bullets(items, styles):
    return ListFlowable(
        [ListItem(Paragraph(item, styles["Normal"]), leftIndent=18) for item in items],
        bulletType="bullet",
        start="circle",
        leftIndent=18,
    )


def pdf_footer(canvas, doc, report: dict[str, Any]) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8.5)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawRightString(7.5 * inch, 0.5 * inch, f"ULTRON 2.7 Phase {report['phase']} Report | Page {doc.page}")
    canvas.restoreState()
