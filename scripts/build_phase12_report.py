from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_12_Report.docx"
BORDER = "B8C2CC"
HEADER_FILL = "F2F4F7"


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(
        doc,
        "Executive Summary",
        [
            "Phase 12 adds a reusable skill registry and local knowledge base to ULTRON 2.7. ULTRON can now list and run built-in skills, ingest safe local documents, search project knowledge, and expose skills, sources, and recent tasks through the web interface.",
            "Skills and knowledge do not bypass safety. Skills validate their own schemas, pass through skill policy, and use the typed tool runtime for OS-facing actions. Knowledge retrieval provides context only; it never becomes execution authority.",
        ],
    )
    bullets(
        doc,
        "Phase 12 Scope",
        [
            "Add a skill registry with names, descriptions, schemas, risk levels, handlers, and examples.",
            "Add built-in notes, reminders, file search, project summary, and daily planning skills.",
            "Add a local knowledge base with safe-root enforcement, text extraction, chunking, metadata, keyword search, and a keyword-only embedding fallback.",
            "Reject obvious secrets during skill input validation and knowledge ingestion.",
            "Add skills and knowledge APIs.",
            "Add UI panels for skills, memory/knowledge, knowledge search, and recent task history.",
            "Add regression tests and architecture documentation.",
        ],
    )
    table(
        doc,
        "What Was Implemented",
        ("Area", "File(s)", "Result"),
        [
            ("Skill registry", "src/ultron27/skills.py", "Schema validation, privacy checks, skill-level policy, built-in skills, examples, and skill-run audit records."),
            ("Knowledge base", "src/ultron27/knowledge.py", "Safe-root document ingestion, Markdown/TXT/PDF/DOCX extraction where available, chunking, metadata, keyword search, and secret rejection."),
            ("Runtime support", "src/ultron27/runtime.py", "Adds handle_tool_call() so skills can submit typed tool calls through validation, policy, executor, and audit."),
            ("Local API", "src/ultron27/web_server.py", "Adds /api/skills, /api/skills/run, /api/knowledge/ingest, and /api/knowledge/search."),
            ("Interface", "web/index.html, web/styles.css, web/app.js", "Adds skills, memory/knowledge, local search, and recent-task panels."),
            ("Launcher", "scripts/run_phase12_skills_ui.py", "Adds a Phase 12 launcher for the web interface."),
            ("Docs", "README.md, docs/architecture.md, docs/phase12_skills_knowledge.md", "Documents the skill/knowledge boundary and APIs."),
            ("Tests", "tests/test_pipeline.py", "Expands to 67 tests covering skills, knowledge, privacy, policy, and endpoints."),
        ],
    )
    table(
        doc,
        "Safety Boundary",
        ("Boundary", "Rule", "Why It Matters"),
        [
            ("Skill schema", "Inputs must match each skill schema.", "Prevents ambiguous or unexpected skill payloads."),
            ("Skill policy", "Medium and high-risk skills require confirmation.", "Keeps reusable skills under the same permission model."),
            ("Tool runtime", "OS-facing skills call handle_tool_call().", "Preserves tool validation, policy, executor restrictions, and audit logging."),
            ("Knowledge paths", "Ingested files must stay inside safe roots.", "Prevents accidental indexing of private filesystem areas."),
            ("Knowledge content", "Obvious secrets are rejected.", "Prevents storing passwords, tokens, API keys, and credentials."),
            ("Retrieval", "Search results are context only.", "Knowledge cannot command ULTRON to execute tasks."),
        ],
    )
    table(
        doc,
        "Verification Results",
        ("Check", "Command", "Result"),
        [
            ("Unit tests", "$env:PYTHONPATH='src'; python -m unittest discover -s tests", "67 tests passed."),
            ("Compile check", "python -m py_compile Phase 12 modules and scripts", "Phase 12 Python modules compiled successfully."),
            ("Frontend syntax", "bundled node --check web/app.js", "Frontend JavaScript parsed successfully."),
            ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 0.9965, tool 0.9965, argument exact match 0.9579, confirmation 1.0."),
            ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases, tool accuracy 1.0, policy action accuracy 1.0."),
            ("API smoke", "Local HTTP checks", "Skills, knowledge ingest/search, and status panels returned expected payloads."),
        ],
    )
    bullets(
        doc,
        "Recommended Next Step",
        [
            "Add optional embeddings with a local provider and keep keyword search as fallback.",
            "Add source-linked answer generation that cites local chunks.",
            "Add a user-reviewed knowledge source manager for deleting and re-ingesting documents.",
        ],
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 12 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 12 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    p = doc.add_paragraph()
    r = p.add_run("Skills and local knowledge base")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)

    table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", "Phase 12: Skills and Local Knowledge"),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", "Implemented and verified"),
        ],
        header=False,
        widths=(Inches(1.55), Inches(4.95)),
    )


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(
    doc: Document,
    heading: str,
    columns: tuple[str, ...],
    rows: list[tuple[str, ...]],
    header: bool = True,
    widths: tuple[Inches, ...] | None = None,
) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        if len(all_rows[0]) == 3:
            widths = (Inches(1.35), Inches(2.35), Inches(2.8))
        else:
            widths = tuple(Inches(6.5 / len(all_rows[0])) for _ in all_rows[0])
    set_table_width(t, widths)
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0, fill=HEADER_FILL if header and r_idx == 0 else None)


def set_table_width(table_obj, widths) -> None:
    tbl_pr = table_obj._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width.twips for width in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table_obj._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width.twips))
        grid.append(col)


def set_cell(cell, bold: bool = False, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)

    if fill:
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill)

    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(9.1)
            run.bold = bold
            if bold:
                run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


if __name__ == "__main__":
    main()
