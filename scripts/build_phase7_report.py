from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_7_Report.docx"
BORDER = "B8C2CC"
HEADER_FILL = "F2F4F7"


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(
        doc,
        "Executive Summary",
        [
            "Phase 7 gives ULTRON 2.7 a local web interface instead of a CLI-only experience. The interface presents a cyberpunk black-and-green scene with a Three.js plasma sphere, subtle particles, live subtitles, typed command input, and explicit visual states for listening, thinking, speaking, and idle.",
            "The UI connects to ULTRON's existing brain/runtime through a local API. It does not bypass schema validation, policy, execution restrictions, or audit logging.",
        ],
    )
    bullets(
        doc,
        "Phase 7 Scope",
        [
            "Create a real local web interface backed by the ULTRON runtime.",
            "Build a responsive Three.js sphere scene with green plasma lines, eclipse rings, particles, glow, and state transitions.",
            "Add subtitle rendering and a subtitle on/off toggle.",
            "Add typed command input with Send button and Enter-key behavior.",
            "Add local API endpoints for commands, status, memory, subtitles, and development visual state control.",
            "Update tests, README, architecture notes, and implementation documentation.",
        ],
    )
    table(
        doc,
        "What Was Implemented",
        ("Area", "File(s)", "Result"),
        [
            ("Local API", "src/ultron27/web_server.py", "Serves the static UI and exposes command, status, memory, subtitle, and state endpoints."),
            ("Launcher", "scripts/run_phase7_ui.py, pyproject.toml", "Adds a script launcher and the ultron-ui package entry point."),
            ("Frontend shell", "web/index.html, web/styles.css", "Adds the full-screen black-and-green interface, subtitle panel, controls, and command bar."),
            ("3D scene", "web/app.js, web/vendor/three.module.min.js", "Adds the Three.js sphere, green particles, rings, plasma arcs, and state transitions."),
            ("Smoke check", "scripts/phase7_visual_smoke.mjs", "Adds a Playwright-based smoke path that captures screenshots when a browser binary is available."),
            ("Docs", "README.md, docs/architecture.md, docs/phase7_visual_interface.md", "Documents the UI, API endpoints, safety boundary, and run instructions."),
            ("Tests", "tests/test_pipeline.py", "Adds Phase 7 API, subtitle, static asset, and command-processing coverage."),
        ],
    )
    table(
        doc,
        "Visual State Design",
        ("State", "Animation", "User Impression"),
        [
            ("Listening", "Sphere contracts slightly, glow tightens, breathing ring calms, particle motion slows.", "ULTRON is focused on the user's next command."),
            ("Thinking", "Green orbit rings spin around the sphere with crescent-like eclipse motion.", "ULTRON is analyzing and planning the request."),
            ("Speaking", "Sphere expands, plasma lines brighten, outer glow intensifies, particles push outward.", "ULTRON is delivering a result."),
            ("Idle", "Sphere remains alive with restrained motion and a stable field.", "The system is online and ready."),
        ],
    )
    table(
        doc,
        "Runtime Safety Boundary",
        ("Layer", "Responsibility", "Safety Property"),
        [
            ("Web UI", "Collect typed commands and display state/subtitles.", "No shell access and no direct executor calls."),
            ("Local API", "Translate HTTP requests into brain calls and response snapshots.", "Only exposes typed endpoints."),
            ("Brain", "Plan and execute task steps through the assistant runtime.", "Cannot skip planner, validator, policy, or executor."),
            ("Runtime", "Validate tool calls, apply policy, execute allowed tasks, and audit.", "High-risk and blocked actions remain gated."),
        ],
    )
    table(
        doc,
        "Verification Results",
        ("Check", "Command", "Result"),
        [
            ("Unit tests", "$env:PYTHONPATH='src'; python -m unittest discover -s tests", "35 tests passed."),
            ("Compile check", "python -m py_compile src/ultron27/web_server.py scripts/run_phase7_ui.py", "Backend and launcher compiled successfully."),
            ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 1.0, tool 1.0, argument exact match 0.9614, confirmation 1.0."),
            ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases, tool accuracy 1.0, policy action accuracy 1.0."),
            ("Browser smoke", "node scripts/phase7_visual_smoke.mjs", "Script is present and degrades cleanly when the Playwright browser binary is unavailable in the sandbox."),
        ],
    )
    bullets(
        doc,
        "Recommended Next Step",
        [
            "Move to Phase 8 by adding voice input and voice output as adapters around the same local API and brain/runtime.",
            "Map voice activity to the existing listening, thinking, and speaking visual states.",
            "Keep confirmation prompts visible in the interface before enabling broader real-world task execution.",
        ],
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 7 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 7 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    p = doc.add_paragraph()
    r = p.add_run("Cyberpunk visual interface and local brain API")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)

    table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", "Phase 7: Visual Interface"),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", "Implemented and verified"),
        ],
        header=False,
        widths=(Inches(1.55), Inches(4.95)),
    )


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(
    doc: Document,
    heading: str,
    columns: tuple[str, ...],
    rows: list[tuple[str, ...]],
    header: bool = True,
    widths: tuple[Inches, ...] | None = None,
) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        if len(all_rows[0]) == 3:
            widths = (Inches(1.45), Inches(2.35), Inches(2.7))
        else:
            widths = tuple(Inches(6.5 / len(all_rows[0])) for _ in all_rows[0])
    set_table_width(t, widths)
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0, fill=HEADER_FILL if header and r_idx == 0 else None)


def set_table_width(table_obj, widths) -> None:
    tbl_pr = table_obj._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width.twips for width in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table_obj._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width.twips))
        grid.append(col)


def set_cell(cell, bold: bool = False, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)

    if fill:
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill)

    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(9.3)
            run.bold = bold
            if bold:
                run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


if __name__ == "__main__":
    main()
