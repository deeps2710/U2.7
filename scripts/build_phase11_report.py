from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_11_Report.docx"
BORDER = "B8C2CC"
HEADER_FILL = "F2F4F7"


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(
        doc,
        "Executive Summary",
        [
            "Phase 11 adds a real Windows automation executor pack while keeping ULTRON 2.7 inside the existing typed-tool safety architecture. The assistant can now front practical laptop actions such as app launch, safe-root file work, notes, reminders, clipboard, screenshots, and optional volume/brightness adapters.",
            "The LLM and planner still do not receive raw shell access. Every action remains user request -> brain/planner -> typed tool call -> schema validation -> policy gate -> executor adapter -> audit log -> response.",
        ],
    )
    bullets(
        doc,
        "Phase 11 Scope",
        [
            "Add a Windows automation adapter module behind the executor.",
            "Restrict app launch to an approved alias registry.",
            "Block explicit file and folder paths outside configured safe roots.",
            "Require confirmation for medium-risk and high-risk actions.",
            "Keep destructive actions non-destructive and not implemented.",
            "Add a UI confirmation modal and confirmation-cancel API.",
            "Add a recent-audit endpoint for runtime inspection.",
            "Add tests and docs for the Windows pack.",
        ],
    )
    table(
        doc,
        "What Was Implemented",
        ("Area", "File(s)", "Result"),
        [
            ("Windows adapter", "src/ultron27/windows_executor.py", "Typed helpers for app aliases, path opening, clipboard, screenshots, optional volume, optional brightness, and safe target validation."),
            ("Executor", "src/ultron27/executor.py", "Routes Windows tasks through the adapter, blocks unsafe explicit paths, handles medium/destructive tools conservatively, and preserves dry-run behavior."),
            ("Policy", "src/ultron27/policy.py, src/ultron27/models.py", "Adds permission levels and requires confirmation for medium-risk and high-risk commands."),
            ("Planner", "src/ultron27/planner.py", "Routes terminal commands to the medium-risk terminal tool instead of generic app launch."),
            ("Audit API", "src/ultron27/audit.py, src/ultron27/web_server.py", "Adds recent audit reading plus GET /api/audit/recent."),
            ("Visual UI", "web/index.html, web/styles.css, web/app.js", "Adds a confirmation modal for risky typed and voice actions, plus cancellation that clears pending voice confirmation."),
            ("Config and launcher", "ultron.config.example.json, scripts/run_phase11_windows_ui.py", "Documents common Windows aliases and adds a Phase 11 launch script."),
            ("Tests", "tests/test_pipeline.py", "Expands to 60 tests covering safe roots, alias resolution, confirmation blocking, destructive non-execution, and audit endpoints."),
        ],
    )
    table(
        doc,
        "Permission Behavior",
        ("Level", "Meaning", "Phase 11 Handling"),
        [
            ("low_risk_allowed", "Validated low-risk tool.", "Allowed after schema validation and policy approval."),
            ("medium_risk_confirmation", "Tool can affect local state or open powerful surfaces.", "Paused until explicit confirmation."),
            ("high_risk_confirmation", "Tool has elevated risk.", "Paused until explicit confirmation."),
            ("destructive_blocked_or_not_implemented", "Delete, move, rename, script, email, shutdown, restart.", "Confirmation-gated, then executor returns not_implemented rather than destructive behavior."),
            ("blocked", "Invalid, unsafe, or unknown execution request.", "Rejected before executor side effects."),
        ],
    )
    table(
        doc,
        "Verification Results",
        ("Check", "Command", "Result"),
        [
            ("Unit tests", "$env:PYTHONPATH='src'; python -m unittest discover -s tests", "60 tests passed."),
            ("Compile check", "python -m py_compile src/ultron27/windows_executor.py src/ultron27/executor.py src/ultron27/policy.py src/ultron27/web_server.py scripts/run_phase11_windows_ui.py", "Phase 11 Python modules compiled successfully."),
            ("Frontend syntax", "bundled node --check web/app.js", "Frontend JavaScript parsed successfully."),
            ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 0.9965, tool 0.9965, argument exact match 0.9579, confirmation 1.0. The intentional miss is Terminal now routing to the safer medium-risk terminal tool."),
            ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases, tool accuracy 1.0, policy action accuracy 1.0."),
            ("API smoke", "Local HTTP checks", "Status, command confirmation, and audit recent routes returned expected payloads."),
        ],
    )
    bullets(
        doc,
        "Recommended Next Step",
        [
            "Install optional Windows packages for real volume and brightness control if those features are needed on the target laptop.",
            "Add a small UI audit viewer panel on top of /api/audit/recent.",
            "Build a signed or user-reviewed app alias editor instead of editing JSON by hand.",
        ],
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 11 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 11 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    p = doc.add_paragraph()
    r = p.add_run("Windows automation executor pack")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)

    table(
        doc,
        "",
        ("Project", "ULTRON 2.7"),
        [
            ("Phase", "Phase 11: Windows Automation Executor Pack"),
            ("Workspace", "Local workspace checkout: U2.7"),
            ("Status", "Implemented and verified"),
        ],
        header=False,
        widths=(Inches(1.55), Inches(4.95)),
    )


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(
    doc: Document,
    heading: str,
    columns: tuple[str, ...],
    rows: list[tuple[str, ...]],
    header: bool = True,
    widths: tuple[Inches, ...] | None = None,
) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        if len(all_rows[0]) == 3:
            widths = (Inches(1.35), Inches(2.35), Inches(2.8))
        else:
            widths = tuple(Inches(6.5 / len(all_rows[0])) for _ in all_rows[0])
    set_table_width(t, widths)
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0, fill=HEADER_FILL if header and r_idx == 0 else None)


def set_table_width(table_obj, widths) -> None:
    tbl_pr = table_obj._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width.twips for width in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table_obj._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width.twips))
        grid.append(col)


def set_cell(cell, bold: bool = False, fill: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)

    if fill:
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill)

    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(9.1)
            run.bold = bold
            if bold:
                run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


if __name__ == "__main__":
    main()
