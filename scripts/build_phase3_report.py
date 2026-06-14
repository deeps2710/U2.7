from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_3_Report.docx"
BORDER = "B8C2CC"


def main() -> None:
    doc = Document()
    configure(doc)
    title(doc)
    section(doc, "Executive Summary", [
        "Phase 3 adds dataset quality and safety evaluation tooling. ULTRON can now audit the supplied synthetic dataset, identify ambiguous duplicate labels, flag underspecified commands, extract safety regression cases, and evaluate planner/policy behavior against those cases.",
        "This phase strengthens the project before LLM or voice integration by making data quality and safety measurable."
    ])
    bullets(doc, "Phase 3 Scope", [
        "Analyze the full 2,500-example dataset.",
        "Detect duplicate utterances with conflicting labels.",
        "Flag underspecified search/open-file examples.",
        "Extract high-risk and confirmation-gated safety cases.",
        "Evaluate policy behavior on the safety regression set.",
        "Document dataset cleanup recommendations."
    ])
    table(doc, "What Was Implemented", ("Area", "File(s)", "Result"), [
        ("Dataset quality module", "src/ultron27/dataset_quality.py", "Reusable analyzer and safety-case extraction."),
        ("Quality script", "scripts/analyze_dataset_quality.py", "Writes JSON report and safety regression JSONL."),
        ("Safety evaluator", "scripts/evaluate_safety_regression.py", "Checks expected tools and policy actions."),
        ("Generated artifacts", "docs/phase3_dataset_quality_report.json, data/regression/safety_cases.jsonl", "Concrete audit output for future phases."),
        ("Tests", "tests/test_pipeline.py", "Expanded to 17 regression tests.")
    ])
    table(doc, "Audit Metrics", ("Metric", "Value", "Meaning"), [
        ("Total examples", "2500", "Full supplied synthetic dataset size."),
        ("Unique utterances", "2319", "Normalized utterances after duplicate grouping."),
        ("Duplicate groups", "101", "Utterances appearing more than once."),
        ("Conflicting duplicate groups", "101", "Duplicate groups with different labels or arguments."),
        ("Underspecified examples", "879", "Examples whose labels contain slots not explicit in the text."),
        ("Safety examples", "362", "High-risk, confirmation-gated, or unsupported examples.")
    ])
    table(doc, "Verification Results", ("Check", "Command", "Result"), [
        ("Unit tests", "python -m unittest discover -s tests", "17 tests passed."),
        ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 1.0, tool 1.0, argument exact match 0.9614, confirmation 1.0."),
        ("Quality audit", "python scripts/analyze_dataset_quality.py", "Report and safety cases generated."),
        ("Safety regression", "python scripts/evaluate_safety_regression.py", "362 cases, tool accuracy 1.0, policy action accuracy 1.0.")
    ])
    bullets(doc, "Recommended Next Step", [
        "Deduplicate conflicting synthetic rows.",
        "Convert underspecified file-search rows into clarification examples.",
        "Add real usage logs with corrected transcripts.",
        "Add Hinglish, Hindi, and STT-error variants.",
        "Keep safety cases as a permanent regression gate."
    ])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 12, "1F4D78")]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 3 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ULTRON 2.7 Phase 3 Execution Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    p = doc.add_paragraph()
    r = p.add_run("Dataset quality and safety regression layer")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(14)
    table(doc, "", ("Project", "ULTRON 2.7"), [
        ("Phase", "Phase 3: Dataset Quality And Safety Evaluation"),
        ("Workspace", "Local workspace checkout: U2.7"),
        ("Status", "Implemented and verified")
    ], header=False)


def section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in paragraphs:
        doc.add_paragraph(item)


def bullets(doc: Document, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc: Document, heading: str, columns: tuple[str, ...], rows: list[tuple[str, ...]], header: bool = True) -> None:
    if heading:
        doc.add_heading(heading, level=1)
    all_rows = [columns, *rows]
    t = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for r_idx, row_values in enumerate(all_rows):
        for c_idx, value in enumerate(row_values):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell(cell, bold=header and r_idx == 0)


def set_cell(cell, bold: bool = False) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders") or OxmlElement("w:tcBorders")
    if borders.getparent() is None:
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        border = borders.find(qn(f"w:{edge}")) or OxmlElement(f"w:{edge}")
        if border.getparent() is None:
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), BORDER)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(9.5)
            run.bold = bold


if __name__ == "__main__":
    main()
