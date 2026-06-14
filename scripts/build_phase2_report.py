from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ULTRON_2.7_Phase_2_Report.docx"
WORKSPACE_LABEL = "Local workspace checkout: U2.7"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
BORDER = "B8C2CC"


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_title(doc)
    add_summary(doc)
    add_scope(doc)
    add_components(doc)
    add_safety(doc)
    add_verification(doc)
    add_next(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.name)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ULTRON 2.7 Phase 2 Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_title(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("ULTRON 2.7 Phase 2 Execution Report")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Safe low-risk laptop execution layer")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    table = doc.add_table(rows=4, cols=2)
    style_table(table, [2200, 7160], fill_header=False)
    for row, values in zip(
        table.rows,
        [
            ("Project", "ULTRON 2.7"),
            ("Phase", "Phase 2: Safe Low-Risk Execution"),
            ("Workspace", WORKSPACE_LABEL),
            ("Status", "Implemented and verified"),
        ],
    ):
        set_row(row, values)
        row.cells[0].paragraphs[0].runs[0].bold = True


def add_summary(doc: Document) -> None:
    doc.add_heading("Executive Summary", level=1)
    para(doc, "Phase 2 upgrades ULTRON from a dry-run-only prototype into a controlled laptop assistant that can perform selected low-risk actions for real.")
    para(doc, "The implementation adds safe-root file boundaries, application aliases, real note/reminder/timer persistence, safe file search/opening, clipboard helpers, screenshot support, and stronger Windows console output handling.")


def add_scope(doc: Document) -> None:
    doc.add_heading("Phase 2 Scope", level=1)
    bullets(doc, [
        "Keep the Phase 1 planner, validator, policy, and audit flow intact.",
        "Allow real execution only for low-risk tools.",
        "Restrict file and folder tools to configured safe roots.",
        "Add config fields for app aliases, safe roots, and screenshot output.",
        "Keep destructive actions non-destructive even after confirmation.",
        "Add tests for safe-root behavior and real low-risk writes.",
    ])


def add_components(doc: Document) -> None:
    doc.add_heading("What Was Implemented", level=1)
    table = doc.add_table(rows=1, cols=3)
    style_table(table, [2300, 3600, 3460])
    set_row(table.rows[0], ("Area", "File(s)", "Result"), header=True)
    rows = [
        ("Configuration", "src/ultron27/config.py", "Added safe_roots, app_aliases, and screenshot_dir."),
        ("Executor", "src/ultron27/executor.py", "Added real handlers for low-risk laptop actions."),
        ("Planner", "src/ultron27/planner.py", "Added regex plans for notes, reminders, timers, screenshot, clipboard, and folders."),
        ("Tool schema", "src/ultron27/tools.py", "Added clipboard tool schemas."),
        ("CLI", "src/ultron27/cli.py", "Passes config to executor and emits Windows-safe JSON."),
        ("Docs", "README.md, docs/architecture.md, docs/phase2_execution.md", "Documents Phase 2 behavior and guardrails."),
        ("Tests", "tests/test_pipeline.py", "Expanded to 15 regression tests."),
    ]
    for item in rows:
        set_row(table.add_row(), item)


def add_safety(doc: Document) -> None:
    doc.add_heading("Safety Boundary", level=1)
    table = doc.add_table(rows=1, cols=3)
    style_table(table, [2200, 3500, 3660])
    set_row(table.rows[0], ("Category", "Tools", "Behavior"), header=True)
    rows = [
        ("Real low-risk", "notes, reminders, timers, search/open safe files, app launch", "Allowed after validation and policy allow."),
        ("Optional low-risk", "screenshot, clipboard", "Runs only when OS/library support exists."),
        ("Protected high-risk", "delete file, send email, run script, shutdown, restart", "Policy may confirm, but executor remains not_implemented."),
        ("File boundary", "search/open file/folder", "Limited to configured safe_roots."),
    ]
    for item in rows:
        set_row(table.add_row(), item)


def add_verification(doc: Document) -> None:
    doc.add_heading("Verification Results", level=1)
    table = doc.add_table(rows=1, cols=3)
    style_table(table, [2500, 3400, 3460])
    set_row(table.rows[0], ("Check", "Command", "Result"), header=True)
    rows = [
        ("Unit tests", "python -m unittest discover -s tests", "15 tests passed."),
        ("Dataset evaluation", "python scripts/evaluate_dataset.py --split test", "Intent 1.0, tool 1.0, argument exact match 0.9614, confirmation 1.0."),
        ("Screenshot dry-run", 'python -m ultron27 "take a screenshot" --no-audit', "Valid low-risk dry-run response."),
        ("Real note write", 'python -m ultron27 "create note phase two" --execute --workspace . --no-audit', "Created notes/phase_two.md."),
        ("Confirmed delete", 'python -m ultron27 "delete project_report.txt" --yes --execute --no-audit', "Returned not_implemented, so no destructive action occurred."),
    ]
    for item in rows:
        set_row(table.add_row(), item)


def add_next(doc: Document) -> None:
    doc.add_heading("Recommended Next Step", level=1)
    para(doc, "Phase 3 should focus on dataset quality and evaluation: deduplicate conflicting synthetic rows, add real usage logs, add Hinglish/Hindi/STT-error variants, and create a separate safety regression split.")


def para(doc: Document, text: str) -> None:
    doc.add_paragraph(text).paragraph_format.space_after = Pt(6)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet").paragraph_format.space_after = Pt(4)


def style_table(table, widths: list[int], fill_header: bool = True) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[col_index])
            set_cell_margins(cell)
            set_cell_border(cell)
            if fill_header and row_index == 0:
                shade_cell(cell, "F2F4F7")


def set_row(row, values, header: bool = False) -> None:
    for cell, value in zip(row.cells, values):
        cell.text = str(value)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if header:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in [("top", 80), ("start", 120), ("bottom", 80), ("end", 120)]:
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), BORDER)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


if __name__ == "__main__":
    main()

